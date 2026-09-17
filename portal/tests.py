from io import BytesIO
from pathlib import Path
import tempfile
import uuid

from django.contrib.auth import get_user_model
from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import course_workflow_views
from .course_services import (
    day_one_list_filename,
    day_one_students_upload_filename,
)
from .identity import format_emirates_id, masked_emirates_id
from .models import (
    ActivityLog,
    Camp,
    Course,
    CourseInstructor,
    CourseInstructorInventoryUsage,
    CourseSession,
    CourseSessionProposal,
    Instructor,
    InstructorInventoryBalance,
    InstructorInventoryMovement,
    InstructorRole,
    InventoryItem,
    Registration,
    SourceFile,
    SourceRecord,
    Student,
    Team,
    TrainingRecord,
)


@override_settings(
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    DEFAULT_FROM_EMAIL="test@iqarus.example",
)
class SimpleCourseWorkflowTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="lead-admin",
            password="test-password",
            is_staff=True,
        )
        self.instructor_user = User.objects.create_user(
            username="course-instructor",
            email="instructor@example.com",
            password="test-password",
        )
        lead_role = InstructorRole.objects.create(name="LI")
        self.instructor = Instructor.objects.create(
            user=self.instructor_user,
            name_english="COURSE INSTRUCTOR",
            email="instructor@example.com",
        )
        self.instructor.roles.add(lead_role)
        self.unassigned_user = User.objects.create_user(
            username="unassigned-instructor",
            email="unassigned@example.com",
            password="test-password",
        )
        instructor_role = InstructorRole.objects.create(name="INS")
        self.unassigned_instructor = Instructor.objects.create(
            user=self.unassigned_user,
            name_english="UNASSIGNED INSTRUCTOR",
            email="unassigned@example.com",
        )
        self.unassigned_instructor.roles.add(instructor_role)
        self.course = Course.objects.create(
            code="ASM",
            title_english="ASM - TCCC",
        )
        self.camp = Camp.objects.create(name="Test Camp")
        self.tempdir = tempfile.TemporaryDirectory()
        self.original_roster_root = course_workflow_views.ROSTER_ROOT
        course_workflow_views.ROSTER_ROOT = Path(self.tempdir.name)

    def tearDown(self):
        course_workflow_views.ROSTER_ROOT = self.original_roster_root
        self.tempdir.cleanup()

    def _create_course(self, instructor_ids=None):
        self.client.force_login(self.admin)
        instructor_ids = instructor_ids or [self.instructor.id]
        response = self.client.post(
            reverse("course_create"),
            {
                "course": self.course.id,
                "camp": self.camp.id,
                "start_date": "2026-09-01",
                "end_date": "2026-09-02",
                "capacity": "20",
                "instructor_student_ratio": "10",
                "service_branch": "PG",
                "instructors": instructor_ids,
                "poc_name": "Camp Contact",
                "poc_contact_number": "+971500000000",
                "poc_location_url": "https://example.com/location",
            },
        )
        self.assertEqual(response.status_code, 302)
        return CourseSession.objects.get(course=self.course)

    def _registration(
        self,
        session,
        *,
        eid="784198012345678",
        name_english="TEST STUDENT",
        name_arabic="طالب تجريبي",
        email="student@example.com",
    ):
        return Registration.objects.create(
            requested_session=session,
            submitted_name_english=name_english,
            submitted_name_arabic=name_arabic,
            eid_raw=eid,
            eid_normalized=eid,
            email_raw=email,
            submitted_unit="Unit",
            status=Registration.Status.PENDING,
        )

    def _history_source_file(self):
        return SourceFile.objects.create(
            filename="historical-evidence.xlsx",
            sha256=uuid.uuid4().hex.ljust(64, "0"),
            records_found=1,
            imported=True,
        )

    def _history_record(
        self,
        source_file,
        *,
        record_type,
        eid,
        raw_payload,
        import_status=SourceRecord.ImportStatus.STAGED,
        linked_instructor=None,
        linked_registration=None,
        linked_training_record=None,
    ):
        return SourceRecord.objects.create(
            source_record_id=uuid.uuid4(),
            source_file=source_file,
            worksheet_name="History",
            source_row_number=SourceRecord.objects.count() + 2,
            record_type=record_type,
            raw_payload=raw_payload,
            derived_search_values={
                "normalized_emirates_id": eid,
            },
            import_status=import_status,
            linked_instructor=linked_instructor,
            linked_registration=linked_registration,
            linked_training_record=linked_training_record,
        )

    def _history_proposal(
        self,
        *records,
        start_date,
        course=None,
        instructor_values=None,
    ):
        proposal = CourseSessionProposal.objects.create(
            candidate_id=uuid.uuid4(),
            proposal_source=(
                CourseSessionProposal.ProposalSource.REGISTRATION
            ),
            proposed_course=course,
            proposed_camp=self.camp,
            start_date=start_date,
            end_date=start_date,
            camp_values=[self.camp.name],
            unit_values=["PG"],
            instructor_values=instructor_values or [],
            evidence_records=len(records),
            student_id_values=len(records),
        )
        proposal.source_records.add(*records)
        return proposal

    def test_emirates_id_display_always_uses_dashes(self):
        raw = "784198012345678"
        expected = "784-1980-1234567-8"
        self.assertEqual(format_emirates_id(raw), expected)
        self.assertEqual(format_emirates_id(expected), expected)
        self.assertEqual(
            masked_emirates_id(raw),
            "•••-••••-••••567-8",
        )

    def test_create_course_uses_assignment_focused_layout(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("course_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Course details")
        self.assertContains(response, "Service Branch")
        self.assertContains(response, "Point of Contact")
        self.assertContains(response, "Assigned instructors")
        self.assertContains(response, 'id="instructor-search"')
        self.assertContains(response, 'id="selected-instructor-count"')
        self.assertNotContains(response, "Lead Instructor")

    def test_course_creation_emails_assignment_and_creates_qr_token(self):
        session = self._create_course()
        assignment = CourseInstructor.objects.get(session=session)
        self.assertEqual(
            assignment.assignment_role,
            CourseInstructor.AssignmentRole.INSTRUCTOR,
        )
        self.assertIsNotNone(assignment.notified_at)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Camp Contact", mail.outbox[0].body)
        self.assertIn("Service Branch: PG", mail.outbox[0].body)
        self.assertIn("https://example.com/location", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["instructor@example.com"])
        self.assertFalse(
            CourseInstructor.objects.filter(
                session=session,
                instructor=self.unassigned_instructor,
            ).exists()
        )
        self.assertTrue(session.public_registration_token)
        self.assertFalse(session.registration_published)

        self.client.force_login(self.unassigned_user)
        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 403)

    def test_course_can_have_multiple_selected_instructors(self):
        session = self._create_course(
            [self.instructor.id, self.unassigned_instructor.id]
        )
        self.assertEqual(
            CourseInstructor.objects.filter(session=session).count(),
            2,
        )
        self.assertEqual(len(mail.outbox), 2)
        self.assertEqual(
            {message.to[0] for message in mail.outbox},
            {"instructor@example.com", "unassigned@example.com"},
        )

    def test_other_service_branch_requires_and_saves_custom_name(self):
        self.client.force_login(self.admin)
        course_data = {
            "course": self.course.id,
            "camp": self.camp.id,
            "start_date": "2026-09-01",
            "end_date": "2026-09-02",
            "capacity": "20",
            "instructor_student_ratio": "10",
            "service_branch": "OTHER",
            "instructors": [self.instructor.id],
            "poc_name": "Camp Contact",
            "poc_contact_number": "+971500000000",
            "poc_location_url": "https://example.com/location",
        }
        response = self.client.post(reverse("course_create"), course_data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter the service branch.")

        course_data["service_branch_other"] = "Coast Guard"
        response = self.client.post(reverse("course_create"), course_data)
        self.assertEqual(response.status_code, 302)
        session = CourseSession.objects.get(course=self.course)
        self.assertEqual(session.service_branch, "OTHER")
        self.assertEqual(session.service_branch_other, "Coast Guard")
        self.assertEqual(session.service_branch_name, "Coast Guard")
        self.assertIn("Service Branch: Coast Guard", mail.outbox[0].body)

    def test_authorised_user_edits_course_and_only_new_assignment_is_emailed(self):
        session = self._create_course()
        assignment = CourseInstructor.objects.get(session=session)
        assignment.acknowledged_at = timezone.now()
        assignment.save(update_fields=["acknowledged_at", "updated_at"])
        second_camp = Camp.objects.create(name="Second Camp")
        mail.outbox.clear()

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("course_edit", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Save Course Changes")

        response = self.client.post(
            reverse("course_edit", args=[session.public_id]),
            {
                "course": self.course.id,
                "camp": second_camp.id,
                "start_date": "2026-09-03",
                "end_date": "2026-09-04",
                "capacity": "30",
                "instructor_student_ratio": "15",
                "service_branch": "NAVY",
                "instructors": [self.unassigned_instructor.id],
                "poc_name": "Updated Contact",
                "poc_contact_number": "+971511111111",
                "poc_location_url": "https://example.com/updated-location",
            },
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.camp, second_camp)
        self.assertEqual(session.capacity, 30)
        self.assertEqual(session.service_branch, "NAVY")
        self.assertEqual(session.poc_name, "Updated Contact")
        self.assertTrue(session.reference_code.endswith("SECOND-CAMP"))
        self.assertFalse(
            CourseInstructor.objects.filter(
                session=session,
                instructor=self.instructor,
            ).exists()
        )
        new_assignment = CourseInstructor.objects.get(
            session=session,
            instructor=self.unassigned_instructor,
        )
        self.assertIsNone(new_assignment.acknowledged_at)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["unassigned@example.com"])
        self.assertIn("Updated Contact", mail.outbox[0].body)
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.UPDATE,
            object_id=str(session.public_id),
            description="Course details or instructor assignments updated.",
        )
        self.assertEqual(
            activity.details["added_instructor_ids"],
            [self.unassigned_instructor.id],
        )
        self.assertEqual(
            activity.details["removed_instructor_ids"],
            [self.instructor.id],
        )

    def test_course_cancel_preserves_course_and_writes_activity_log(self):
        session = self._create_course()
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_cancel", args=[session.public_id]),
            {"reason": "Camp requested cancellation"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.status, CourseSession.Status.CANCELLED)
        self.assertFalse(session.registration_published)
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.CANCEL,
            object_id=str(session.public_id),
        )
        self.assertEqual(
            activity.details["reason"],
            "Camp requested cancellation",
        )

    def test_empty_course_can_be_deleted_but_activity_remains(self):
        session = self._create_course()
        public_id = str(session.public_id)
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_delete", args=[session.public_id]),
            {"confirm": "delete", "reason": "Created by mistake"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(CourseSession.objects.filter(public_id=public_id).exists())
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.DELETE,
            object_id=public_id,
        )
        self.assertEqual(activity.details["reason"], "Created by mistake")

    def test_course_with_registration_cannot_be_deleted(self):
        session = self._create_course()
        self._registration(session)
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_delete", args=[session.public_id]),
            {"confirm": "delete"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(CourseSession.objects.filter(pk=session.pk).exists())
        self.assertFalse(
            ActivityLog.objects.filter(
                action=ActivityLog.Action.DELETE,
                object_id=str(session.public_id),
            ).exists()
        )

    def test_inventory_is_held_and_reported_per_instructor(self):
        item = InventoryItem.objects.create(
            name="Training Kit",
            category=InventoryItem.Category.CONSUMABLE,
            unit="kit",
            quantity_on_hand=20,
        )
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("inventory"),
            {
                "action": "allocate",
                "instructor_id": self.unassigned_instructor.id,
                "item_id": item.id,
                "movement_type": "issue",
                "quantity": "8",
            },
        )
        self.assertEqual(response.status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.quantity_on_hand, 12)
        balance = InstructorInventoryBalance.objects.get(
            instructor=self.unassigned_instructor,
            item=item,
        )
        self.assertEqual(balance.quantity_on_hand, 8)

        session = self._create_course(
            [self.instructor.id, self.unassigned_instructor.id]
        )
        self.client.force_login(self.unassigned_user)
        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.unassigned_instructor.id,
                f"used_{item.id}": "5",
                f"consumed_{item.id}": "2",
                f"deteriorated_{item.id}": "1",
                f"inventory_notes_{item.id}": "Course use",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 5)
        usage = CourseInstructorInventoryUsage.objects.get(
            session=session,
            instructor=self.unassigned_instructor,
            item=item,
        )
        self.assertEqual(usage.quantity_used, 5)
        self.assertEqual(usage.quantity_consumed, 2)
        self.assertEqual(usage.quantity_deteriorated, 1)
        movement = InstructorInventoryMovement.objects.filter(
            session=session,
            instructor=self.unassigned_instructor,
            item=item,
        ).latest("created_at")
        self.assertEqual(movement.quantity_change, -3)
        self.assertEqual(movement.balance_after, 5)

        # Saving the same report again must not deduct the balance twice.
        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.unassigned_instructor.id,
                f"used_{item.id}": "5",
                f"consumed_{item.id}": "2",
                f"deteriorated_{item.id}": "1",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 5)

        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {"instructor_id": self.instructor.id},
        )
        self.assertEqual(response.status_code, 403)

    def test_material_used_layout_saves_one_used_up_value_per_instructor(self):
        item = InventoryItem.objects.create(
            name="Pressure Bandage",
            category=InventoryItem.Category.CONSUMABLE,
            unit="each",
            quantity_on_hand=0,
        )
        balance = InstructorInventoryBalance.objects.create(
            instructor=self.instructor,
            item=item,
            quantity_on_hand=20,
            updated_by=self.admin,
        )
        session = self._create_course()
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Material Used")
        self.assertContains(response, "Balance Before Course")
        self.assertContains(response, "Used Up in Course")
        self.assertContains(response, "Balance After Course")
        self.assertContains(response, f'name="used_up_{item.id}"')
        self.assertNotContains(response, f'name="consumed_{item.id}"')
        self.assertNotContains(response, f'name="deteriorated_{item.id}"')

        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.instructor.id,
                f"used_up_{item.id}": "5",
                f"inventory_notes_{item.id}": "5 damaged",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 15)
        usage = CourseInstructorInventoryUsage.objects.get(
            session=session,
            instructor=self.instructor,
            item=item,
        )
        self.assertEqual(usage.quantity_used, 5)
        self.assertEqual(usage.quantity_consumed, 5)
        self.assertEqual(usage.quantity_deteriorated, 0)

    def test_inventory_records_use_tabs_filters_and_25_row_pages(self):
        for number in range(26):
            item = InventoryItem.objects.create(
                name=f"Filter Item {number:02d}",
                category=InventoryItem.Category.EQUIPMENT,
                unit="item",
                quantity_on_hand=0,
            )
            InstructorInventoryBalance.objects.create(
                instructor=self.instructor,
                item=item,
                quantity_on_hand=number + 1,
                updated_by=self.admin,
            )
            InstructorInventoryMovement.objects.create(
                instructor=self.instructor,
                item=item,
                movement_type=InstructorInventoryMovement.MovementType.ISSUE,
                quantity_change=number + 1,
                balance_after=number + 1,
                recorded_by=self.admin,
            )

        self.client.force_login(self.admin)
        response = self.client.get(reverse("inventory"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "portal/js/scroll_position.")
        self.assertContains(response, 'id="inventory-balances-panel"')
        self.assertNotContains(response, 'id="inventory-history-panel"')
        self.assertEqual(len(response.context["balance_page"]), 25)
        self.assertTrue(response.context["balance_page"].has_next())

        response = self.client.get(
            reverse("inventory"),
            {"view": "balances", "q": "Filter Item 25"},
        )
        self.assertEqual(response.context["balance_page"].paginator.count, 1)

        response = self.client.get(
            reverse("inventory"),
            {"view": "history"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="inventory-history-panel"')
        self.assertNotContains(response, 'id="inventory-balances-panel"')
        self.assertEqual(len(response.context["movement_page"]), 25)
        self.assertTrue(response.context["movement_page"].has_next())

        response = self.client.get(
            reverse("inventory"),
            {"view": "history", "page": "2"},
        )
        self.assertEqual(len(response.context["movement_page"]), 1)

    def test_assigned_instructor_acknowledges_and_controls_registration(self):
        session = self._create_course()
        self.client.force_login(self.instructor_user)
        response = self.client.post(
            reverse("course_acknowledge", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 302)
        assignment = CourseInstructor.objects.get(session=session)
        self.assertIsNotNone(assignment.acknowledged_at)
        response = self.client.post(
            reverse("registration_qr", args=[session.public_id]),
            {"action": "open"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertTrue(session.registration_published)
        response = self.client.post(
            reverse("registration_qr", args=[session.public_id]),
            {"action": "close"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertFalse(session.registration_published)

    def test_single_roster_saves_independent_tags_and_exact_template(self):
        session = self._create_course()
        registration = self._registration(session)
        self._registration(
            session,
            eid="784198112345679",
            name_english="NOT SELECTED STUDENT",
            name_arabic="طالب غير مختار",
            email="not-selected@example.com",
        )
        self.client.force_login(self.instructor_user)
        response = self.client.post(
            reverse("course_roster_save_download", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"hp_{registration.public_id}": "yes",
                f"ttt_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "passed",
                f"remarks_{registration.public_id}": "Clear practical result",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            "attachment; filename*=UTF-8''01-02_September_COURSE_Test_Camp.xlsx",
        )
        registration.refresh_from_db()
        self.assertTrue(registration.selected_for_roster)
        self.assertTrue(registration.is_hp)
        self.assertTrue(registration.is_ttt)
        self.assertEqual(registration.assessment_status, "passed")
        self.assertEqual(TrainingRecord.objects.count(), 1)
        record = TrainingRecord.objects.get()
        self.assertTrue(record.is_hp)
        self.assertTrue(record.is_ttt)

        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["List 1-1"])
        sheet = workbook.active
        self.assertEqual(
            [sheet.cell(3, column).value for column in range(1, 9)],
            [
                "#",
                "Email",
                "EID / رقم الهوية",
                "اسم الطالب",
                "Student Name",
                "Day 1",
                "Day 2",
                "Status",
            ],
        )
        self.assertEqual(sheet["B4"].value, "student@example.com")
        self.assertEqual(sheet["C4"].value, "784-1980-1234567-8")
        self.assertEqual(sheet["D4"].value, "طالب تجريبي")
        self.assertEqual(sheet["E4"].value, "TEST STUDENT")
        self.assertIsNone(sheet["H4"].value)
        self.assertEqual(sheet["E5"].value, "Status: Pass, Fail, D (Duplicate), TTT (Train the Trainer)")
        self.assertIsNone(sheet["A5"].value)
        self.assertTrue(str(sheet.print_area).endswith("$H$7"))
        self.assertNotIn(
            0,
            [
                cell.value
                for row in sheet.iter_rows(min_row=1, max_row=7)
                for cell in row
            ],
        )
        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertContains(response, "784-1980-1234567-8")
        self.assertNotContains(response, "Review required")

    def test_workspace_bulk_controls_are_in_their_table_columns(self):
        session = self._create_course()
        self._registration(session)
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="column-control">Select')
        self.assertContains(response, 'id="hp-all"')
        self.assertContains(response, 'id="ttt-all"')
        self.assertContains(response, 'id="pass-all"')
        self.assertContains(response, 'id="fail-all"')
        self.assertContains(response, 'class="student-row')
        self.assertContains(response, ".roster{min-width:0;table-layout:fixed}")
        self.assertNotContains(response, 'class="toolbar"')

    def test_dashboard_and_course_pages_use_reference_layout(self):
        session = self._create_course()
        registration = self._registration(session)
        registration.selected_for_roster = True
        registration.assessment_status = "passed"
        registration.is_hp = True
        registration.save(
            update_fields=[
                "selected_for_roster",
                "assessment_status",
                "is_hp",
                "updated_at",
            ]
        )
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="dashboard-filters"')
        self.assertContains(response, 'class="summary-counter-card"', count=3)
        self.assertContains(response, 'class="dashboard-chart"')
        self.assertContains(response, 'class="dashboard-module-grid"')
        self.assertContains(response, 'name="all_history"')
        self.assertEqual(response.context["totals"]["courses"], 1)
        self.assertEqual(response.context["totals"]["enrolled"], 1)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["hp"], 1)

        response = self.client.get(reverse("my_courses"), {"q": "COURSE INSTRUCTOR"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Search by course, camp, or instructor....")
        self.assertContains(response, "Course Name")
        self.assertContains(response, "Date(s)")
        self.assertContains(response, "Enrolled 1")
        self.assertContains(response, "Passed 1")

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="course-stage-flow"')
        self.assertContains(response, "Generate Stamp List")
        self.assertContains(response, "Upload Stamp List")
        self.assertContains(response, "Generate Front-End")
        self.assertContains(response, "Upload Front-End")

    def test_dashboard_combines_history_and_current_without_double_counting(self):
        current_session = self._create_course()
        current_student = Student.objects.create(
            eid="784198012345678",
            name_english="CURRENT STUDENT",
        )
        current_registration = self._registration(current_session)
        current_registration.student = current_student
        current_registration.selected_for_roster = True
        current_registration.assessment_status = "passed"
        current_registration.is_hp = True
        current_registration.save()
        TrainingRecord.objects.create(
            student=current_student,
            session=current_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
            raw_payload={
                "registration_public_id": str(current_registration.public_id)
            },
        )
        pending_registration = self._registration(
            current_session,
            eid="784198112345678",
            name_english="PENDING STUDENT",
            email="pending@example.com",
        )
        pending_registration.selected_for_roster = True
        pending_registration.save(update_fields=["selected_for_roster", "updated_at"])

        historical_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date="2026-02-01",
            end_date="2026-02-02",
            status=CourseSession.Status.COMPLETED,
        )
        CourseInstructor.objects.create(
            session=historical_session,
            instructor=self.instructor,
        )
        historical_student = Student.objects.create(
            eid="784198212345678",
            name_english="HISTORICAL STUDENT",
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.FAIL,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_ttt=True,
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
            duplicate_flag=True,
        )
        Registration.objects.create(
            requested_session=historical_session,
            submitted_name_english="REJECTED STUDENT",
            status=Registration.Status.REJECTED,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["courses"], 2)
        self.assertEqual(response.context["totals"]["enrolled"], 3)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["failed"], 1)
        self.assertEqual(response.context["totals"]["rejected"], 1)
        self.assertEqual(response.context["totals"]["hp"], 1)
        self.assertEqual(response.context["totals"]["ttt"], 1)
        february = response.context["chart_rows"][1]
        september = response.context["chart_rows"][8]
        self.assertEqual(february["enrolled"], 1)
        self.assertEqual(february["failed"], 1)
        self.assertEqual(february["rejected"], 1)
        self.assertEqual(september["enrolled"], 2)
        self.assertEqual(september["passed"], 1)

    def test_dashboard_keeps_date_range_and_offers_all_history(self):
        current_session = self._create_course()
        current_student = Student.objects.create(
            eid="784198312345678",
            name_english="CURRENT YEAR STUDENT",
        )
        TrainingRecord.objects.create(
            student=current_student,
            session=current_session,
            result=TrainingRecord.Result.FAIL,
        )
        historical_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date="2024-03-10",
            end_date="2024-03-11",
            status=CourseSession.Status.COMPLETED,
        )
        historical_student = Student.objects.create(
            eid="784198412345678",
            name_english="OLDER HISTORY STUDENT",
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            result=TrainingRecord.Result.PASS,
        )

        self.client.force_login(self.admin)
        current_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertFalse(current_response.context["all_history"])
        self.assertEqual(current_response.context["totals"]["courses"], 1)
        self.assertEqual(current_response.context["totals"]["enrolled"], 1)
        self.assertEqual(current_response.context["totals"]["failed"], 1)

        all_response = self.client.get(
            reverse("dashboard"),
            {
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
                "all_history": "1",
            },
        )
        self.assertTrue(all_response.context["all_history"])
        self.assertContains(all_response, 'name="all_history" value="1" checked')
        self.assertEqual(all_response.context["totals"]["courses"], 2)
        self.assertEqual(all_response.context["totals"]["enrolled"], 2)
        self.assertEqual(all_response.context["totals"]["passed"], 1)
        self.assertEqual(all_response.context["totals"]["failed"], 1)

        historical_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2024-01-01", "date_to": "2024-12-31"},
        )
        self.assertFalse(historical_response.context["all_history"])
        self.assertEqual(historical_response.context["totals"]["courses"], 1)
        self.assertEqual(historical_response.context["totals"]["enrolled"], 1)
        self.assertEqual(historical_response.context["totals"]["passed"], 1)
        self.assertEqual(historical_response.context["totals"]["failed"], 0)

    def test_dashboard_includes_deduplicated_historical_evidence(self):
        session = self._create_course()
        current_student = Student.objects.create(
            eid="784198012345678",
            name_english="CURRENT STUDENT",
        )
        TrainingRecord.objects.create(
            student=current_student,
            session=session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
        )

        source_file = self._history_source_file()
        historical_registration = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198112345678",
            raw_payload={
                "date_التاريخ": {
                    "value": "2026-08-03T00:00:00",
                    "excel_type": "datetime",
                },
                "camp_المعسكر": "Historical Camp",
                "unit_location_الوحدة_الموقع": "PG",
                "instructor": "Course Instructor",
            },
            import_status=SourceRecord.ImportStatus.CONVERTED,
        )
        self._history_proposal(
            historical_registration,
            start_date="2026-08-03",
            course=self.course,
        )

        attendance = self._history_record(
            source_file,
            record_type="course_attendance_evidence",
            eid="784198212345678",
            raw_payload={"session1": "YES", "session2": "YES"},
        )
        result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198212345678",
            raw_payload={"comment": "PASS"},
        )
        self._history_proposal(
            attendance,
            result,
            start_date="2026-08-19",
            course=self.course,
        )

        self._history_record(
            source_file,
            record_type="ttt_candidate_evidence",
            eid="784198312345678",
            raw_payload={
                "training_date": "16-Jul_26",
                "camp_name": "Historical Camp",
                "unit": "PG",
                "instructor_1": "Course Instructor",
            },
        )
        ignored_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198412345678",
            raw_payload={"comment": "FAIL"},
            import_status=SourceRecord.ImportStatus.IGNORED,
        )
        self._history_proposal(
            ignored_result,
            start_date="2026-08-20",
            course=self.course,
        )
        Student.objects.create(
            eid="784198512345678",
            name_english="PROFILE WITHOUT PARTICIPATION",
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )

        totals = response.context["totals"]
        self.assertEqual(totals["courses"], 1)
        self.assertEqual(totals["enrolled"], 3)
        self.assertEqual(totals["passed"], 2)
        self.assertEqual(totals["failed"], 0)
        self.assertEqual(totals["hp"], 1)
        self.assertEqual(totals["ttt"], 1)
        self.assertEqual(response.context["chart_rows"][7]["enrolled"], 2)
        self.assertEqual(response.context["chart_rows"][7]["passed"], 1)

        course_response = self.client.get(
            reverse("dashboard"),
            {
                "course": str(self.course.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(course_response.context["totals"]["enrolled"], 3)

        other_course = Course.objects.create(
            code="OTHER",
            title_english="Other Course",
        )
        other_response = self.client.get(
            reverse("dashboard"),
            {
                "course": str(other_course.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(other_response.context["totals"]["enrolled"], 0)

    def test_dashboard_historical_evidence_respects_dates_and_all_history(self):
        source_file = self._history_source_file()
        current_year = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198612345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        self._history_proposal(
            current_year,
            start_date="2026-08-03",
            course=self.course,
        )
        old_attendance = self._history_record(
            source_file,
            record_type="course_attendance_evidence",
            eid="784198712345678",
            raw_payload={"session1": "YES", "session2": "YES"},
        )
        old_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198712345678",
            raw_payload={"comment": "PASS"},
        )
        self._history_proposal(
            old_attendance,
            old_result,
            start_date="2024-03-10",
            course=self.course,
        )
        self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198812345678",
            raw_payload={"camp_المعسكر": "Undated Camp"},
        )

        self.client.force_login(self.admin)
        current_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(current_response.context["totals"]["enrolled"], 1)
        self.assertEqual(current_response.context["totals"]["passed"], 0)

        all_response = self.client.get(
            reverse("dashboard"),
            {
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
                "all_history": "1",
            },
        )
        self.assertEqual(all_response.context["totals"]["enrolled"], 3)
        self.assertEqual(all_response.context["totals"]["passed"], 1)
        self.assertEqual(all_response.context["chart_rows"][2]["enrolled"], 1)
        self.assertEqual(all_response.context["chart_rows"][2]["passed"], 1)

        old_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2024-01-01", "date_to": "2024-12-31"},
        )
        self.assertEqual(old_response.context["totals"]["enrolled"], 1)
        self.assertEqual(old_response.context["totals"]["passed"], 1)

    def test_dashboard_canonical_training_precedes_source_evidence(self):
        session = self._create_course()
        student = Student.objects.create(
            eid="784198912345678",
            name_english="CANONICAL STUDENT",
        )
        training = TrainingRecord.objects.create(
            student=student,
            session=session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
        )
        source_file = self._history_source_file()
        source_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid=student.eid,
            raw_payload={"comment": "FAIL"},
            linked_training_record=training,
        )
        proposal = self._history_proposal(
            source_result,
            start_date="2026-09-01",
            course=self.course,
        )
        proposal.approved_session = session
        proposal.save(update_fields=["approved_session", "updated_at"])

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 1)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["failed"], 0)

    def test_dashboard_linked_rejection_is_not_counted_twice(self):
        session = self._create_course()
        source_file = self._history_source_file()
        rejected_registration = Registration.objects.create(
            requested_session=session,
            submitted_name_english="REJECTED STUDENT",
            eid_raw="784-1994-1234567-8",
            eid_normalized="784199412345678",
            status=Registration.Status.REJECTED,
            selected_for_roster=True,
            source_file=source_file,
        )
        source = self._history_record(
            source_file,
            record_type="registration_submission",
            eid=rejected_registration.eid_normalized,
            raw_payload={"date_\u0627\u0644\u062a\u0627\u0631\u064a\u062e": "2026-09-01"},
            linked_registration=rejected_registration,
        )
        self._history_proposal(
            source,
            start_date="2026-09-01",
            course=self.course,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 0)
        self.assertEqual(response.context["totals"]["rejected"], 1)

    def test_dashboard_history_survives_approval_and_pending_duplicate_review(self):
        session = self._create_course()
        source_file = self._history_source_file()
        first_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199312345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        repeated_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199312345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        self._history_proposal(
            first_submission,
            repeated_submission,
            start_date="2026-08-03",
            course=self.course,
        )

        approved_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199412345678",
            raw_payload={"date_التاريخ": "2026-09-01"},
        )
        approved_proposal = self._history_proposal(
            approved_submission,
            start_date="2026-09-01",
            course=self.course,
        )
        approved_proposal.approved_session = session
        approved_proposal.proposal_status = (
            CourseSessionProposal.ProposalStatus.APPROVED
        )
        approved_proposal.save(
            update_fields=[
                "approved_session",
                "proposal_status",
                "updated_at",
            ]
        )

        cancelled_student = Student.objects.create(
            eid="784199512345678",
            name_english="CANCELLED CANONICAL ROW",
        )
        cancelled_training = TrainingRecord.objects.create(
            student=cancelled_student,
            session=session,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.CANCELLED,
        )
        retained_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid=cancelled_student.eid,
            raw_payload={"comment": "PASS"},
            linked_training_record=cancelled_training,
        )
        self._history_proposal(
            retained_result,
            start_date="2026-09-01",
            course=self.course,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 4)
        self.assertEqual(response.context["totals"]["passed"], 1)

    def test_dashboard_historical_team_and_instructor_scope_are_trusted(self):
        team_a = Team.objects.create(name="Historical Team A")
        team_b = Team.objects.create(name="Historical Team B")
        self.unassigned_instructor.team = team_a
        self.unassigned_instructor.save(update_fields=["team", "updated_at"])
        self.instructor.team = team_b
        self.instructor.save(update_fields=["team", "updated_at"])

        source_file = self._history_source_file()
        own = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199012345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
            linked_instructor=self.unassigned_instructor,
        )
        other = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199112345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
            linked_instructor=self.instructor,
        )
        unresolved = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199212345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        for record in (own, other, unresolved):
            self._history_proposal(
                record,
                start_date="2026-08-03",
                course=self.course,
            )

        self.client.force_login(self.admin)
        team_response = self.client.get(
            reverse("dashboard"),
            {
                "team": str(team_a.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(team_response.context["totals"]["enrolled"], 1)

        all_teams_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(all_teams_response.context["totals"]["enrolled"], 3)

        self.client.force_login(self.unassigned_user)
        instructor_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(instructor_response.context["totals"]["enrolled"], 1)

    def test_existing_student_and_previous_same_course_require_review(self):
        session = self._create_course()
        master = Student.objects.create(
            eid="784198012345678",
            name_english="TEST STUDENT",
            name_arabic="طالب تجريبي",
            email="student@example.com",
        )
        earlier_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            reference_code="ASM-PREVIOUS",
            start_date="2026-08-01",
            end_date="2026-08-02",
            status=CourseSession.Status.COMPLETED,
        )
        TrainingRecord.objects.create(
            student=master,
            session=earlier_session,
            result=TrainingRecord.Result.PASS,
        )
        registration = self._registration(session)
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertContains(response, "Existing master student")
        self.assertContains(response, "Previously attended this course")
        self.assertContains(response, "Review required")

        response = self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "pending",
            },
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertFalse(registration.selected_for_roster)

        response = self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"duplicate_decision_{registration.public_id}": "reviewed_allow",
                f"duplicate_notes_{registration.public_id}": "Identity checked",
                f"result_{registration.public_id}": "pending",
            },
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertTrue(registration.selected_for_roster)
        self.assertEqual(
            registration.duplicate_review_status,
            Registration.DuplicateReviewStatus.ALLOWED,
        )
        self.assertEqual(registration.duplicate_review_notes, "Identity checked")
        self.assertEqual(len(registration.duplicate_review_fingerprint), 64)
        self.assertEqual(registration.duplicate_reviewed_by, self.instructor_user)

        self.client.force_login(self.admin)
        response = self.client.get(reverse("data_quality") + "?state=allowed")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Identity checked")
        self.assertContains(response, "Previously attended this course")

    def test_external_upload_contains_only_passed_selected_student(self):
        session = self._create_course()
        registration = self._registration(session)
        self.client.force_login(self.instructor_user)
        self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "passed",
            },
        )
        response = self.client.get(
            reverse("external_upload_xlsx", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            "attachment; filename*=UTF-8''"
            "01-02_September_COURSE_Test_Camp_Students%20upload.xlsx",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["obs", "TCCC ASM"])
        sheet = workbook["TCCC ASM"]
        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 6)],
            ["email", "emirates_id", "full_name", "session1", "session2"],
        )
        self.assertEqual(sheet["C2"].value, "طالب تجريبي")
        self.assertEqual(sheet["B2"].value, "784-1980-1234567-8")
        self.assertEqual(sheet["D2"].value, "YES")
        self.assertEqual(sheet["E2"].value, "YES")

    def test_day_one_filenames_match_supplied_examples(self):
        session = self._create_course()
        self.instructor.name_english = "Vasilios Example"
        self.instructor.save(update_fields=["name_english", "updated_at"])
        self.unassigned_instructor.name_english = "Abdulsattar Dalgamuni"
        self.unassigned_instructor.save(
            update_fields=["name_english", "updated_at"]
        )
        CourseInstructor.objects.create(
            session=session,
            instructor=self.unassigned_instructor,
        )
        session.start_date = timezone.datetime(2026, 8, 19).date()
        session.end_date = timezone.datetime(2026, 8, 20).date()
        self.camp.name = "Specialized Reserve Nahel PG"
        self.camp.save(update_fields=["name", "updated_at"])

        expected_base = (
            "19-20_August_Vasilios-Abdulsattar_"
            "Specialized_Reserve_Nahel_PG"
        )
        self.assertEqual(day_one_list_filename(session), expected_base + ".xlsx")
        self.assertEqual(
            day_one_students_upload_filename(session),
            expected_base + "_Students upload.xlsx",
        )

    def test_course_lists_use_compact_human_references(self):
        session = self._create_course()
        self.client.force_login(self.admin)

        response = self.client.get(reverse("my_courses"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "course-row")
        self.assertContains(response, "View Course")
        self.assertContains(response, "Camp")
        self.assertContains(response, "Test Camp")
        self.assertContains(response, 'class="instructor-tag"')
        self.assertContains(response, 'class="course-instructors-label">Instructors</strong>')
        self.assertContains(response, "COURSE INSTRUCTOR")
        self.assertNotContains(response, "ASM · TCCC · 01 SEP 2026")
        self.assertNotContains(response, session.reference_code)

        response = self.client.get(reverse("courses_sessions"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "MASTER DIRECTORY")
        self.assertContains(response, "ASM - TCCC")
        self.assertContains(response, "Course Directory")

    def test_reports_offer_focused_one_row_views(self):
        session = self._create_course()
        registration = self._registration(session)
        registration.selected_for_roster = True
        registration.assessment_status = "passed"
        registration.is_hp = True
        registration.save(
            update_fields=[
                "selected_for_roster",
                "assessment_status",
                "is_hp",
                "updated_at",
            ]
        )
        item = InventoryItem.objects.create(
            name="Report Kit",
            category=InventoryItem.Category.EQUIPMENT,
            unit="kit",
        )
        CourseInstructorInventoryUsage.objects.create(
            session=session,
            instructor=self.instructor,
            item=item,
            quantity_used=2,
        )
        self.client.force_login(self.admin)

        for view_name, expected in [
            ("progress", "Course progress"),
            ("outcomes", "Student outcomes"),
            ("documents", "Document status"),
            ("instructors", "COURSE INSTRUCTOR"),
            ("inventory", "Report Kit"),
            ("activity", "Activity log"),
        ]:
            response = self.client.get(
                reverse("reports") + f"?view={view_name}"
            )
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, expected)

    def test_student_directory_print_uses_day_one_browser_sheet(self):
        student = Student.objects.create(
            eid="784198012345678",
            name_english="PRINT TEST STUDENT",
            name_arabic="طالب للطباعة",
            email="print@example.com",
        )
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("students_print"),
            {
                "mode": "selected",
                "student_ids": [student.pk],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Student Directory List")
        self.assertContains(response, "EID / رقم الهوية")
        self.assertContains(response, "Day 1")
        self.assertContains(response, "Day 2")
        self.assertContains(response, "Status")
        self.assertContains(response, "PRINT TEST STUDENT")
        self.assertTrue(
            ActivityLog.objects.filter(
                action=ActivityLog.Action.PRINT,
                object_type="Student",
            ).exists()
        )

    def test_admin_navigation_exposes_course_directory(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Directory", count=2)
        self.assertContains(response, "iqarus-logo-transparent")

        self.client.force_login(self.unassigned_user)
        response = self.client.get(reverse("courses_sessions"))
        self.assertEqual(response.status_code, 403)
